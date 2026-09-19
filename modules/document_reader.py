"""
modules/document_reader.py  --  file bytes  ->  plain text.

This is the ONLY place that knows about file formats. Everything after this
works on text, so adding a new format means adding one function here.

    read_document(data, filename) -> ReadResult

Supported:  .txt   .pdf (text layer)   .pdf/.png/.jpg (scanned -> OCR)
            .docx  .xlsx

A reader never crashes the pipeline. If a file cannot be read, it returns
ok=False and a `problem` code, so the caller can send it to human review.
    problem codes: corrupt_file | empty_document | ocr_unavailable |
                   unsupported_type | read_error
"""
import io
import os
from dataclasses import dataclass, field
from typing import Optional


@dataclass
class ReadResult:
    text: str = ""
    method: str = "unknown"          # text | pdf_text | ocr | docx | xlsx
    ok: bool = True
    problem: Optional[str] = None    # see problem codes above
    detail: str = ""                 # human-readable explanation
    ocr_confidence: Optional[float] = None   # 0..1, only for OCR
    warnings: list = field(default_factory=list)
    images: list = field(default_factory=list, repr=False)   # page images of a SCAN (for the AI vision step)


def _fail(method, problem, detail):
    return ReadResult(text="", method=method, ok=False, problem=problem, detail=detail)


# ---------------------------------------------------------------- plain text
def _read_txt(data: bytes) -> ReadResult:
    text = data.decode("utf-8", errors="replace")
    if not text.strip():
        return _fail("text", "empty_document", "text file is empty")
    return ReadResult(text=text, method="text")


# ----------------------------------------------------------------------- PDF
MIN_TEXT_CHARS = 20  # fewer characters than this => treat the PDF as a scan


def _read_pdf(data: bytes) -> ReadResult:
    try:
        import pdfplumber
    except ImportError:
        return _fail("pdf_text", "read_error", "pdfplumber is not installed (pip install pdfplumber)")

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if not pdf.pages:
                return _fail("pdf_text", "empty_document", "PDF has no pages")
            pages = [_page_to_text(p) for p in pdf.pages]
            text = "\n".join(pages)
            if len(text.strip()) >= MIN_TEXT_CHARS:
                return ReadResult(text=text, method="pdf_text")
            # No text layer -> it is a scan. Render pages to images and OCR them.
            images = [p.to_image(resolution=300).original for p in pdf.pages]
    except Exception as e:  # corrupt / truncated / encrypted PDF
        return _fail("pdf_text", "corrupt_file", f"cannot open PDF: {type(e).__name__}: {e}")

    return _ocr_images(images)


def _page_to_text(page) -> str:
    """
    Rebuild the page as text that keeps the label / value COLUMNS lined up:

        Shipper          ACME LTD
                         12 MAIN ROAD          <- indented = continuation of Shipper
        Consignee        FOO GMBH

    We work from word positions instead of pdfplumber's layout mode because long
    labels can overlap the value column; character-level layout then interleaves
    the two ("ConsCigEnReIEeX"). Word-level placement keeps them apart.
    """
    words = page.extract_words(use_text_flow=True, x_tolerance=1.5)
    if not words:
        return ""
    widths = sorted((w["x1"] - w["x0"]) / max(len(w["text"]), 1) for w in words)
    cw = widths[len(widths) // 2] or 5.0                     # typical character width
    left = min(w["x0"] for w in words)

    # group words into visual lines (same 'top' within a few points)
    words.sort(key=lambda w: (w["top"], w["x0"]))
    lines, current, top = [], [], None
    for w in words:
        if top is None or abs(w["top"] - top) <= 3:
            current.append(w)
            top = w["top"] if top is None else top
        else:
            lines.append(current)
            current, top = [w], w["top"]
    if current:
        lines.append(current)

    out = []
    for lw in lines:
        lw.sort(key=lambda w: w["x0"])
        line, prev_x1 = "", None
        for w in lw:
            col = int(round((w["x0"] - left) / cw))
            if not line:
                line = " " * col
            elif w["x0"] - prev_x1 <= cw * 1.8:              # normal word gap
                line += " "
            else:                                            # column gap -> keep the columns
                line += " " * max(2, col - len(line))
            line += w["text"]
            prev_x1 = w["x1"]
        out.append(line)
    return "\n".join(out)


# ----------------------------------------------------------------------- OCR
def _ocr_images(images) -> ReadResult:
    """OCR a list of PIL images, and keep the images so the AI vision step can use them too."""
    res = _run_ocr(images)
    res.images = list(images)
    return res


def _run_ocr(images) -> ReadResult:
    """The actual Tesseract call. Returns text + mean word confidence (0..1)."""
    try:
        import pytesseract
        from PIL import ImageOps
        from pytesseract import Output
    except ImportError:
        return _fail("ocr", "ocr_unavailable", "pytesseract is not installed (pip install pytesseract)")

    all_lines, confs = [], []
    try:
        for img in images:
            img = ImageOps.autocontrast(img.convert("L"))  # greyscale + contrast helps Tesseract
            d = pytesseract.image_to_data(img, config="--psm 6", output_type=Output.DICT)
            lines = {}
            for i, word in enumerate(d["text"]):
                if not word.strip():
                    continue
                key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
                lines.setdefault(key, []).append(word)
                c = float(d["conf"][i])
                if c >= 0:
                    confs.append(c)
            all_lines.extend(" ".join(ws) for _, ws in sorted(lines.items()))
    except pytesseract.TesseractNotFoundError:
        return _fail("ocr", "ocr_unavailable",
                     "Tesseract program not found. Install it, or set pytesseract.pytesseract.tesseract_cmd")
    except Exception as e:
        return _fail("ocr", "read_error", f"OCR failed: {type(e).__name__}: {e}")

    text = "\n".join(all_lines)
    if not text.strip():
        return _fail("ocr", "empty_document", "OCR found no text (blank or too-poor scan)")
    mean_conf = (sum(confs) / len(confs) / 100.0) if confs else 0.0
    res = ReadResult(text=text, method="ocr", ocr_confidence=round(mean_conf, 3))
    res.warnings.append("text came from OCR - values may contain reading errors")
    return res


def _read_image(data: bytes) -> ReadResult:
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as e:
        return _fail("ocr", "corrupt_file", f"cannot open image: {e}")
    return _ocr_images([img])


# ---------------------------------------------------------------------- DOCX
def _read_docx(data: bytes) -> ReadResult:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return _fail("docx", "read_error", "python-docx is not installed (pip install python-docx)")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        return _fail("docx", "corrupt_file", f"cannot open DOCX: {type(e).__name__}: {e}")

    out = []
    # Walk the body in order so paragraphs and tables keep their position.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, document).text.strip()
            if t:
                out.append(t)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells, seen = [], set()
                for c in row.cells:            # merged cells repeat -> de-duplicate
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip())
                out.extend(_row_to_lines(cells))
    text = "\n".join(out)
    if not text.strip():
        return _fail("docx", "empty_document", "DOCX has no text")
    return ReadResult(text=text, method="docx")


# ---------------------------------------------------------------------- XLSX
def _read_xlsx(data: bytes) -> ReadResult:
    try:
        import openpyxl
    except ImportError:
        return _fail("xlsx", "read_error", "openpyxl is not installed (pip install openpyxl)")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        return _fail("xlsx", "corrupt_file", f"cannot open XLSX: {type(e).__name__}: {e}")

    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_to_str(v) for v in row]
            cells = [c for c in cells if c != ""]
            if cells:
                out.extend(_row_to_lines(cells))
    text = "\n".join(out)
    if not text.strip():
        return _fail("xlsx", "empty_document", "XLSX has no data")
    return ReadResult(text=text, method="xlsx")


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _row_to_lines(cells):
    """
    Turn one table/sheet row into the same 'Label: value' text a .txt file uses.
    Address parts separated by ' | ' or newlines become indented continuation
    lines, e.g.   Shipper: ACME LTD
                    12 MAIN ROAD; SINGAPORE
    """
    cells = [c for c in cells if c is not None]
    if len(cells) == 1:
        return [cells[0]]
    if len(cells) == 2:
        label, value = cells
        parts = [p.strip() for p in value.replace(" | ", "\n").split("\n") if p.strip()]
        if not parts:
            return [f"{label}:"]
        return [f"{label}: {parts[0]}"] + [f"  {p}" for p in parts[1:]]
    return ["  ".join(c for c in cells if c)]     # wide table row


# ------------------------------------------------------------------ dispatch
def read_document(data: bytes, filename: str) -> ReadResult:
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext in (".txt", ".text", ""):
            return _read_txt(data)
        if ext == ".pdf":
            return _read_pdf(data)
        if ext == ".docx":
            return _read_docx(data)
        if ext in (".xlsx", ".xlsm"):
            return _read_xlsx(data)
        if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            return _read_image(data)
    except Exception as e:  # last-resort safety net
        return _fail("unknown", "read_error", f"{type(e).__name__}: {e}")
    return _fail("unknown", "unsupported_type", f"file type '{ext}' is not supported")